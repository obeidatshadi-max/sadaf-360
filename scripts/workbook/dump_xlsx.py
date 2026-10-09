import sys,json,datetime,openpyxl
from openpyxl.utils.datetime import to_excel
wb=openpyxl.load_workbook(sys.argv[1])
out={}
for ws in wb:
    rows=[]
    for r in ws.iter_rows(min_row=1,max_row=ws.max_row,max_col=ws.max_column):
        row=[]
        for c in r:
            v=c.value
            if isinstance(v,(datetime.datetime,datetime.date)): v=to_excel(v)
            row.append(v)
        rows.append(row)
    out[ws.title]=rows
json.dump(out,open(sys.argv[2],'w'))
print({k:len(v) for k,v in out.items()})
